"""Generate a sample multiple-choice question document (.docx) for testing.

Produces data/nurseflow/sample-quiz-questions.docx — a plain, clearly formatted
quiz that the Studio "Extract from PDF / Word" endpoint can turn into draft cards.

    python scripts/make_sample_quiz_docx.py
"""
import os

from docx import Document
from docx.shared import Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "nurseflow", "sample-quiz-questions.docx")

QUESTIONS = [
    {
        "q": "Before starting prescribed oxygen via nasal cannula, which assessment should the nurse prioritise?",
        "options": [
            "Inspect the nasal passages and assess respiratory status",
            "Offer the patient a glass of water",
            "Place the patient flat in bed",
            "Measure the patient's height",
        ],
        "answer": "A",
        "rationale": "A baseline respiratory assessment and checking the nares support safe application and later reassessment of oxygen therapy.",
        "source": "https://www.brit-thoracic.org.uk/quality-improvement/guidelines/emergency-oxygen/",
    },
    {
        "q": "Which action is most appropriate when a patient shows signs of possible sepsis and clinical deterioration?",
        "options": [
            "Escalate urgently using the local deterioration pathway",
            "Wait until the next routine observations",
            "Give oral fluids and review tomorrow",
            "Document it only after the shift ends",
        ],
        "answer": "A",
        "rationale": "New confusion, fever and increased respiratory rate require timely assessment and escalation under local policy.",
        "source": "https://www.nice.org.uk/guidance/ng51",
    },
    {
        "q": "When should hand hygiene be performed during preparation for an aseptic procedure?",
        "options": [
            "Immediately before the aseptic task",
            "Only after the procedure",
            "At the end of the shift",
            "Only if hands look visibly soiled",
        ],
        "answer": "A",
        "rationale": "Hand hygiene immediately before an aseptic task helps prevent microorganisms reaching a vulnerable site.",
        "source": "https://www.who.int/publications/m/item/my-5-moments-for-hand-hygiene",
    },
    {
        "q": "A medication label does not match the electronic prescription. What should the nurse do first?",
        "options": [
            "Stop and clarify the discrepancy before administration",
            "Administer the labelled dose to avoid delay",
            "Ask the patient which dose they usually take",
            "Document the dose as given",
        ],
        "answer": "A",
        "rationale": "Do not administer when a discrepancy is present. Pause, verify the order and follow local medicines-management procedures.",
        "source": "https://www.nmc.org.uk/standards/standards-for-post-registration/standards-for-medicines-management/",
    },
    {
        "q": "A patient reports pain rated 7 out of 10. Which action best reflects the nursing process?",
        "options": [
            "Assess the pain, then plan and evaluate interventions against a comfort-function goal",
            "Record the score and take no further action",
            "Delay all care until the pain subsides",
            "Ask a colleague to score the pain instead",
        ],
        "answer": "A",
        "rationale": "Pain assessment should lead to a documented plan and re-evaluation of the patient's comfort-function goal.",
        "source": "https://wtcs.pressbooks.pub/nursingfundamentals/",
    },
    {
        "q": "Which finding most suggests fluid volume deficit in an adult patient?",
        "options": [
            "Dry mucous membranes with decreased urine output",
            "Bounding pulse with peripheral oedema",
            "Weight gain over two days",
            "Crackles at the lung bases",
        ],
        "answer": "A",
        "rationale": "Dry mucous membranes and reduced urine output are classic signs of dehydration and fluid volume deficit.",
        "source": "https://wtcs.pressbooks.pub/nursingfundamentals/",
    },
    {
        "q": "Why is a patient's own account of their sleep important when planning care?",
        "options": [
            "It identifies individual sleep needs and disturbances that guide interventions",
            "It replaces the need for any nursing assessment",
            "Sleep quality cannot be assessed any other way",
            "It is only used for patients with diagnosed insomnia",
        ],
        "answer": "A",
        "rationale": "Sleep is subjective; the patient's report helps the nurse individualise sleep-promotion strategies.",
        "source": "https://wtcs.pressbooks.pub/nursingfundamentals/",
    },
    {
        "q": "A nurse notices a colleague has not performed hand hygiene. What is the most appropriate first step?",
        "options": [
            "Speak to the colleague directly and respectfully about the observation",
            "Ignore it to avoid conflict",
            "Report the colleague to the media",
            "Document it in the patient's notes only",
        ],
        "answer": "A",
        "rationale": "Professional accountability and speaking up for safety start with a direct, respectful conversation under local policy.",
        "source": "https://www.nmc.org.uk/standards/code/",
    },
]


def main() -> None:
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    doc = Document()

    title = doc.add_heading("Sample Nursing Practice Questions", level=0)
    doc.add_paragraph(
        "Original NCLEX-style and NMC-aligned items for testing question extraction. "
        "Each item lists four options, the answer, a rationale and a source."
    )

    for i, item in enumerate(QUESTIONS, start=1):
        doc.add_heading(f"Question {i}", level=2)
        doc.add_paragraph(item["q"])
        for letter, option in zip("ABCD", item["options"]):
            doc.add_paragraph(f"{letter}. {option}")
        p = doc.add_paragraph()
        p.add_run("Answer: ").bold = True
        p.add_run(item["answer"])
        p = doc.add_paragraph()
        p.add_run("Rationale: ").bold = True
        p.add_run(item["rationale"])
        p = doc.add_paragraph()
        p.add_run("Source: ").bold = True
        p.add_run(item["source"]).font.size = Pt(10)

    doc.save(OUT)
    print("Wrote", OUT)


if __name__ == "__main__":
    main()
